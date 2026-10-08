import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from bot.services.rekap_buku_kas import COLUMNS, build_rows, generate_pdf, generate_xlsx, totals
from bot.services.rekap_service import parse_note_text

D = date(2026, 7, 15)


def test_parse_lines_and_types():
    ents, skipped = parse_note_text(
        "beli kopi 150rb\n+ penjualan hari ini 2jt\ngaji barista Rp 1.500.000\nhalo semua", D, "Budi")
    assert [(e.type, e.amount) for e in ents] == [
        ("keluar", 150000), ("masuk", 2_000_000), ("keluar", 1_500_000)]
    assert ents[0].uraian == "beli kopi" and ents[2].kategori == "Gaji"
    assert skipped == ["halo semua"]


def test_explicit_date_overrides_forward_date():
    ents, _ = parse_note_text("listrik 10/07 350000", D)
    assert ents[0].tanggal == date(2026, 7, 10) and ents[0].amount == 350000


def test_header_catatan_grup_dibaca_reconbot():
    # reconbot (catatan_grup.py) mewajibkan header ini; JANGAN diubah
    assert COLUMNS == ["No", "Tanggal", "No. Bukti", "Uraian", "Kategori", "Sumber",
                       "Debit (Masuk)", "Kredit (Keluar)", "Saldo"]


def test_debit_adalah_uang_masuk_kredit_adalah_uang_keluar():
    ents, _ = parse_note_text("+ setoran 1jt\nbeli gula 200rb", D, "Budi")
    rows = build_rows(ents, Decimal(500_000))
    assert [r[8] for r in rows] == [500_000, 1_500_000, 1_300_000]
    assert rows[1][6] == 1_000_000 and rows[1][7] is None      # masuk -> kolom Debit (Masuk)
    assert rows[2][7] == 200_000 and rows[2][6] is None        # keluar -> kolom Kredit (Keluar)
    assert rows[1][2] == "KM-260715-01" and rows[2][2] == "KK-260715-01"
    assert rows[1][5] == "Budi"
    assert totals(rows) == (1_000_000, 200_000, 1_300_000)


def test_excel_debit_kredit_dan_saldo_awal_berupa_angka():
    # reconbot membaca dengan data_only=True: Debit/Kredit dan baris Saldo Awal harus angka biasa
    ents, _ = parse_note_text("+ setoran 1jt\nbeli gula 200rb", D)
    ws = load_workbook(io.BytesIO(generate_xlsx(build_rows(ents, Decimal(500_000))))).active
    hdr = next(r for r in range(1, 12) if ws.cell(row=r, column=3).value == "No. Bukti")
    assert [ws.cell(row=hdr, column=c).value for c in range(1, 10)] == COLUMNS
    assert ws.cell(row=hdr + 1, column=4).value == "Saldo Awal" and ws.cell(row=hdr + 1, column=9).value == 500_000
    assert ws.cell(row=hdr + 2, column=7).value == 1_000_000      # Debit (Masuk)
    assert ws.cell(row=hdr + 3, column=8).value == 200_000        # Kredit (Keluar)


def test_pdf_valid():
    ents, _ = parse_note_text("+ setoran 1jt\nbeli gula 200rb", D)
    assert generate_pdf(build_rows(ents))[:4] == b"%PDF"
