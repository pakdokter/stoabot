"""
Rekap service — parsing catatan forward Telegram + buku kas 9 kolom (XLSX & PDF).

Kolom buku kas:
  1 No | 2 Tanggal | 3 No. Bukti | 4 Uraian | 5 Kategori
  6 Sumber | 7 Debit (Masuk) | 8 Kredit (Keluar) | 9 Saldo
"""
import io
import re
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from typing import Optional

from bot.config import settings
from bot.utils.formatters import fmt_date, fmt_rupiah, parse_amount

COLUMNS = [
    "No", "Tanggal", "No. Bukti", "Uraian", "Kategori",
    "Sumber", "Debit (Masuk)", "Kredit (Keluar)", "Saldo",
]

MASUK_KEYWORDS = (
    "terima", "diterima", "masuk", "setoran", "setor", "penjualan", "jual",
    "omzet", "pendapatan", "pemasukan", "dp", "pelunasan",
)
KATEGORI_KEYWORDS = {
    "Bahan Baku": ("kopi", "susu", "gula", "es batu", "sirup", "bahan", "beras", "telur", "roti", "teh"),
    "Operasional": ("listrik", "air", "wifi", "internet", "pulsa", "gas", "token", "sewa"),
    "Gaji": ("gaji", "upah", "bonus", "lembur"),
    "Peralatan": ("alat", "gelas", "cup", "mesin", "perbaikan", "servis"),
    "Penjualan": ("penjualan", "omzet", "jual"),
}

_AMOUNT_RE = re.compile(
    r"(?:rp\.?\s*)?(\d[\d.,]*\s*(?:juta|jt|ribu|rb)?)(?![\w/])", re.IGNORECASE
)
_DATE_RE = re.compile(r"\b(\d{1,2})[/-](\d{1,2})(?:[/-](\d{2,4}))?\b")


@dataclass
class RekapEntry:
    tanggal: date
    type: str  # masuk | keluar
    amount: float
    uraian: str
    kategori: str = "Umum"
    sumber: str = ""
    file_id: Optional[str] = None
    ocr_raw_text: Optional[str] = None
    ocr_confidence: Optional[float] = None


def guess_type(text: str, sign: str = "") -> str:
    if sign == "+":
        return "masuk"
    if sign == "-":
        return "keluar"
    words = set(re.findall(r"[a-z]+", text.lower()))
    return "masuk" if words & set(MASUK_KEYWORDS) else "keluar"


def guess_category(text: str, tx_type: str) -> str:
    low = text.lower()
    for kategori, kws in KATEGORI_KEYWORDS.items():
        if kategori == "Penjualan" and tx_type != "masuk":
            continue
        if any(re.search(rf"\b{re.escape(k)}\b", low) for k in kws):
            return kategori
    return "Penjualan" if tx_type == "masuk" else "Umum"


def _extract_date(line: str, default: date) -> tuple[str, date]:
    m = _DATE_RE.search(line)
    if not m:
        return line, default
    d, mo = int(m.group(1)), int(m.group(2))
    y = m.group(3)
    year = default.year if not y else (int(y) + 2000 if len(y) == 2 else int(y))
    try:
        found = date(year, mo, d)
    except ValueError:
        return line, default
    return (line[:m.start()] + line[m.end():]).strip(), found


def parse_note_text(text: str, default_date: date, sumber: str = "") -> tuple[list[RekapEntry], list[str]]:
    """Parse teks catatan: satu baris = satu transaksi. Return (entries, baris_diabaikan)."""
    entries: list[RekapEntry] = []
    skipped: list[str] = []
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        line, tgl = _extract_date(line, default_date)
        sign = ""
        if line[:1] in "+-" and len(line) > 1:
            sign, line = line[0], line[1:].strip()

        matches = list(_AMOUNT_RE.finditer(line))
        amount = None
        chosen = None
        for m in reversed(matches):
            val = parse_amount(m.group(1))
            if val and val > 0:
                amount, chosen = val, m
                break
        if amount is None:
            skipped.append(raw.strip())
            continue

        uraian = (line[:chosen.start()] + " " + line[chosen.end():])
        uraian = re.sub(r"\brp\.?\b", "", uraian, flags=re.IGNORECASE)
        uraian = re.sub(r"[\s:=\-–]+$|^[\s:=\-–]+", "", re.sub(r"\s+", " ", uraian)).strip()
        if not uraian:
            skipped.append(raw.strip())
            continue
        tx_type = guess_type(uraian, sign)
        entries.append(RekapEntry(
            tanggal=tgl, type=tx_type, amount=amount, uraian=uraian,
            kategori=guess_category(uraian, tx_type), sumber=sumber,
        ))
    return entries, skipped


def build_rows(entries: list[RekapEntry], saldo_awal: Decimal = Decimal(0)) -> list[list]:
    """Susun baris 9 kolom, urut tanggal. Baris pertama (No kosong) = Saldo Awal."""
    ordered = sorted(entries, key=lambda e: e.tanggal)  # stable: urutan forward dipertahankan
    counters: dict[str, int] = {}
    saldo = Decimal(saldo_awal)
    rows: list[list] = [["", "", "", "Saldo Awal", "", "", None, None, saldo]]
    for i, e in enumerate(ordered, 1):
        prefix = f"{'KM' if e.type == 'masuk' else 'KK'}-{e.tanggal:%y%m%d}"
        counters[prefix] = counters.get(prefix, 0) + 1
        amt = Decimal(str(e.amount))
        debit = amt if e.type == "masuk" else None
        kredit = amt if e.type == "keluar" else None
        saldo += (debit or 0) - (kredit or 0)
        rows.append([
            i, e.tanggal, f"{prefix}-{counters[prefix]:02d}", e.uraian,
            e.kategori, e.sumber, debit, kredit, saldo,
        ])
    return rows


def totals(rows: list[list]) -> tuple[Decimal, Decimal, Decimal]:
    debit = sum((r[6] or 0) for r in rows[1:])
    kredit = sum((r[7] or 0) for r in rows[1:])
    return Decimal(debit), Decimal(kredit), rows[-1][8]


def _periode(rows: list[list]) -> str:
    dates = [r[1] for r in rows[1:]]
    if not dates:
        return ""
    lo, hi = min(dates), max(dates)
    return fmt_date(lo) if lo == hi else f"{fmt_date(lo)} s/d {fmt_date(hi)}"


def generate_xlsx(rows: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Buku Kas"
    ws["A1"] = settings.business_name
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "REKAP BUKU KAS"
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = f"Periode: {_periode(rows)}"

    hdr_row = 5
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=hdr_row, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A1A2E")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    r0 = hdr_row + 1
    for i, row in enumerate(rows):
        r = r0 + i
        for c, val in enumerate(row, 1):
            if isinstance(val, Decimal):
                val = float(val)
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = border
            if c in (7, 8, 9):
                cell.number_format = "#,##0"
            elif c == 2:
                cell.number_format = "DD/MM/YYYY"
            elif c == 4:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 0:
            ws.cell(row=r, column=4).font = Font(bold=True)
        else:
            # Saldo = saldo sebelumnya + debit - kredit (formula, mudah diaudit)
            ws.cell(row=r, column=9, value=f"=I{r-1}+N(G{r})-N(H{r})")

    last = r0 + len(rows) - 1
    tot = last + 1
    ws.cell(row=tot, column=4, value="TOTAL").font = Font(bold=True)
    ws.cell(row=tot, column=7, value=f"=SUM(G{r0+1}:G{last})")
    ws.cell(row=tot, column=8, value=f"=SUM(H{r0+1}:H{last})")
    ws.cell(row=tot, column=9, value=f"=I{last}")
    for c in range(1, 10):
        cell = ws.cell(row=tot, column=c)
        cell.font = Font(bold=True)
        cell.border = border
        if c >= 7:
            cell.number_format = "#,##0"

    for col, w in zip("ABCDEFGHI", (5, 12, 16, 40, 14, 18, 16, 16, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=r0, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf(rows: list[list]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=8)
    title = ParagraphStyle("t", parent=ss["Heading1"], fontSize=14, spaceAfter=2)

    def money(v):
        return "" if v is None else fmt_rupiah(v).replace("Rp", "")

    data = [COLUMNS]
    for r in rows:
        data.append([
            r[0], fmt_date(r[1]) if r[1] else "", r[2],
            Paragraph(str(r[3]), cell), Paragraph(str(r[4]), cell), Paragraph(str(r[5]), cell),
            money(r[6]), money(r[7]), money(r[8]),
        ])
    d, k, s = totals(rows)
    data.append(["", "", "", "TOTAL", "", "", money(d), money(k), money(s)])

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=1 * cm, rightMargin=1 * cm,
        topMargin=1 * cm, bottomMargin=1 * cm,
    )
    widths = [1 * cm, 2.3 * cm, 3.2 * cm, 6.5 * cm, 2.8 * cm, 3.2 * cm, 2.9 * cm, 2.9 * cm, 2.9 * cm]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("ALIGN", (6, 1), (8, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#eeeeee")),
    ]))
    doc.build([
        Paragraph(settings.business_name, title),
        Paragraph(f"REKAP BUKU KAS — {_periode(rows)}", ss["Normal"]),
        Spacer(1, 0.3 * cm),
        table,
    ])
    return buf.getvalue()
