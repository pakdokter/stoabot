import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from bot.services.kas_report import (
    KamusKategori, normalisasi_item,
    kategori_transfer,
    STD_HEADER, KasRow, generate_pdf, generate_xlsx, hitung_saldo, label_kantong,
    map_kategori, pihak, ringkas,
)

D = Decimal


def contoh_rows():
    return [
        KasRow(date(2026, 7, 3), "Uang masuk kantong widia", kredit=D(1_000_000), subjek="-", objek="Kantong Widia"),
        KasRow(date(2026, 7, 4), "Pasar — Bawang Putih", debit=D(-90_000), kategori="Belanja Bahan",
               subjek="Kantong Widia", objek="Pasar"),
        KasRow(date(2026, 7, 5), "Goormet — Dimsum Keju", debit=D(-185_000), subjek="Kantong Widia", objek="Goormet"),
    ]


def test_header_sama_dengan_reconbot():
    assert STD_HEADER == ["Tanggal", "Items / Activities", "Kategori Transaksi", "Debit", "Kredit",
                          "Saldo Kumulatif", "Subjek Transaksi", "Objek Transaksi",
                          "Keterangan Tambahan", "Catatan Pribadi"]


def test_saldo_dan_total_debit_negatif():
    rows = contoh_rows()
    assert hitung_saldo(rows, D(50_000)) == [D(1_050_000), D(960_000), D(775_000)]
    assert ringkas(rows, D(50_000)) == (D(-275_000), D(1_000_000), D(775_000))


def test_layout_excel_persis_reconbot():
    wb = load_workbook(io.BytesIO(generate_xlsx([("Widia Sari", contoh_rows(), D(50_000))])))
    ws = wb["Widia Sari"]
    assert [c.value for c in ws[1]] == STD_HEADER
    # baris 2 = Saldo Awal Bulan (positif -> kolom Kredit), F = saldo awal
    assert (ws["B2"].value, ws["C2"].value, ws["E2"].value, ws["F2"].value) == ("Saldo Awal", "Saldo Awal Bulan", 50000, 50000)
    assert ws["G2"].value == ws["H2"].value == "-"
    # transaksi mulai baris 3, saldo berformula, Debit negatif
    assert ws["F3"].value == "=D3+E3+F2" and ws["E3"].value == 1_000_000
    assert ws["D4"].value == -90_000 and ws["F5"].value == "=D5+E5+F4"
    assert ws["A3"].number_format == "dd/mm/yyyy"
    # satu baris kosong (6), footer 7-10, pemilik kantong 11
    assert all(ws.cell(row=6, column=c).value is None for c in range(1, 11))
    assert ws["B7"].value == "Saldo Awal" and ws["F7"].value == "=F2"
    assert ws["B8"].value == "Total Debit (Uang Keluar)" and ws["D8"].value == "=SUM(D3:D5)"
    assert ws["B9"].value == "Total Kredit (Uang Masuk)" and ws["E9"].value == "=SUM(E3:E5)"
    assert ws["B10"].value == "Saldo Akhir" and ws["F10"].value == "=F2+D8+E9"
    assert (ws["B11"].value, ws["C11"].value) == ("Pemilik Kantong", "Widia Sari")
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:J5"


def test_saldo_awal_negatif_masuk_kolom_debit():
    wb = load_workbook(io.BytesIO(generate_xlsx([("Kiki", [], D(-500))])))
    ws = wb["Kiki"]
    assert ws["D2"].value == -500 and ws["E2"].value is None and ws["F2"].value == -500


def test_banyak_kantong_satu_workbook_dan_nama_sheet_aman():
    wb = load_workbook(io.BytesIO(generate_xlsx([
        ("Widia", contoh_rows(), D(0)), ("Mita/Wandia:*", [], D(0)), ("Widia", [], D(0))])))
    assert wb.sheetnames == ["Widia", "Mita Wandia", "Widia 2"]


def test_pdf_valid():
    pdf = generate_pdf("Widia Sari", contoh_rows(), D(0), "01 Jul 2026 s/d 31 Jul 2026")
    assert pdf[:4] == b"%PDF" and len(pdf) > 1500


def test_kategori_dipetakan_ke_daftar_resmi_atau_kosong():
    assert map_kategori("pasar") == "Belanja Bahan" and map_kategori("Bahan Baku") == "Belanja Bahan"
    assert map_kategori("Umum") == "" and map_kategori(None) == ""


def test_pihak_subjek_objek():
    own = label_kantong("Widia Sari")
    assert pihak("Pasar — Bawang Putih", "keluar", own) == (own, "Pasar")
    assert pihak("SISA UANG", "keluar", own) == (own, "-")
    assert pihak("Widia — Pindah Ke Kantong Widia", "keluar", own) == (own, "Kantong Widia")
    assert pihak("dari kantong mita", "masuk", own) == ("Kantong Mita", own)
    assert pihak("Uang masuk", "masuk", own) == ("-", own)
    assert pihak("Kasir", "masuk", own) == ("Kasir", own)


def test_transfer_ke_kasir_atau_kantong_lain_jadi_transaksi_internal():
    own = label_kantong("Widia Sari")
    assert kategori_transfer("Kasir", own, own) == "Transaksi Internal"          # uang masuk dari Kasir
    assert kategori_transfer("Kantong Mita", own, own) == "Transaksi Internal"   # dari kantong lain
    assert kategori_transfer(own, "Kantong Mita", own) == "Transaksi Internal"   # pindah ke kantong lain
    assert kategori_transfer(own, "Pasar", own) == ""                            # belanja biasa
    assert kategori_transfer("-", own, own) == ""


def test_normalisasi_item_buang_jumlah_dan_ukuran():
    assert normalisasi_item("Primer Raya — MAIZENA TIMBANGAN 500GR x6") == "maizena timbangan"
    assert normalisasi_item("Dinda Frozen Food — Sosis x2.0") == "sosis"
    assert normalisasi_item("Mak opik — Bamer ¼") == "bamer"
    assert normalisasi_item("Primer Raya (6 item)") == "primer raya"


def test_kamus_kategori_hanya_dari_alias_yang_diberikan():
    kk = KamusKategori({"Sosis": "Belanja Bahan", "Thinwall": "Kemasan", "Daging Slice Grade B": "Belanja Bahan"})
    assert kk.cari("Dinda Frozen Food — Sosis x2.0") == "Belanja Bahan"
    assert kk.cari("Primer Raya — Thinwall") == "Kemasan"
    assert kk.cari("Fadhilah — DAGING SLICE GRADE B x2") == "Belanja Bahan"
    assert kk.cari("Primer Raya — Sosis Yomas 1Bks") == "Belanja Bahan"   # frasa utuh di dalam nama item
    assert kk.cari("Fadhilah — Parkir") == ""                               # bukan bahan -> kosong
    assert kk.cari("") == ""


def test_dari_transaksi_pakai_kamus_hanya_untuk_uang_keluar():
    import types
    kk = KamusKategori({"Sosis": "Belanja Bahan"})
    def tx(typ, desc, cat=None):
        return types.SimpleNamespace(transaction_date=date(2026, 7, 3), type=typ, amount=10_000,
                                     description=desc, category=cat, attachments=[])
    from bot.services.kas_report import dari_transaksi
    own = label_kantong("Widia")
    assert dari_transaksi(tx("keluar", "Toko — Sosis"), own, kk).kategori == "Belanja Bahan"
    assert dari_transaksi(tx("keluar", "Toko — Sosis", "Umum"), own, kk).kategori == "Belanja Bahan"
    assert dari_transaksi(tx("masuk", "Sosis"), own, kk).kategori == ""
    assert dari_transaksi(tx("keluar", "Toko — Parkir"), own, kk).kategori == ""


def test_bank_biru_dan_rekening_bank_dikenali_sebagai_sumber(monkeypatch):
    from bot.config import settings
    monkeypatch.setattr(settings, "bank_sumber", "BRI-567(Biz),Jago")
    own = label_kantong("Royyan Paice")
    assert pihak("Bank Biru", "masuk", own) == ("Bank Biru", own)
    assert pihak("BRI-567(Biz)", "masuk", own) == ("BRI-567(Biz)", own)
    assert pihak("Jago", "masuk", own) == ("Jago", own)
    assert pihak("Transfer Owner", "masuk", own) == ("-", own)          # teks bebas tetap tidak ditebak
    for sumber in ("Bank Biru", "Kasir", "BRI-567(Biz)", "Jago", "Kantong Mita"):
        assert kategori_transfer(sumber, own, own) == "Transaksi Internal"
    assert kategori_transfer("Transfer Owner", own, own) == ""


def test_tanpa_bank_sumber_hanya_bank_biru_dan_kasir(monkeypatch):
    from bot.config import settings
    monkeypatch.setattr(settings, "bank_sumber", "")
    own = label_kantong("Mita")
    assert pihak("Jago", "masuk", own) == ("-", own)
    assert pihak("Bank Biru", "masuk", own) == ("Bank Biru", own)


def test_pemilik_yang_bentrok_dengan_sheet_bank_pakai_nama_satu_token(monkeypatch):
    from bot.config import settings
    monkeypatch.setattr(settings, "bank_sumber", "BCA - Royyan")
    assert label_kantong("Royyan Paice") == "Royyan-Kantong"
    assert label_kantong("Mita Wandia") == "Kantong Mita Wandia"          # tidak bentrok: tidak berubah
    wb = load_workbook(io.BytesIO(generate_xlsx([("Royyan Paice", [], D(0)), ("Mita Wandia", [], D(0))])))
    assert wb.sheetnames == ["Royyan-Kantong", "Mita Wandia"]
    own = label_kantong("Mita Wandia")
    assert pihak("dari kantong royyan", "masuk", own) == ("Royyan-Kantong", own)
    assert pihak("Pindah Ke Kantong Royyan", "keluar", own) == (own, "Royyan-Kantong")
    assert pihak("BCA - Royyan", "masuk", label_kantong("Royyan Paice")) == ("BCA - Royyan", "Royyan-Kantong")
    assert kategori_transfer("Royyan-Kantong", own, own) == "Transaksi Internal"
    assert kategori_transfer("BCA - Royyan", "Royyan-Kantong", "Royyan-Kantong") == "Transaksi Internal"


def test_tanpa_bank_sumber_nama_tidak_berubah(monkeypatch):
    from bot.config import settings
    monkeypatch.setattr(settings, "bank_sumber", "")
    assert label_kantong("Royyan Paice") == "Kantong Royyan Paice"
